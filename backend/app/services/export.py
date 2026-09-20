import io
import re


def safe_filename(title: str, max_len: 60 = 60) -> str:
    base = re.sub(r"[^a-zA-Z0-9-_ ]+", "", (title or "research")).strip()
    base = re.sub(r"\s+", "-", base)
    return (base[:max_len] or "research").strip("-")


def clean_inline(text: str) -> str:
    """Strip markdown literals (** bold, `code`, ### headings, [links](url)) to plain text."""
    t = text or ""
    t = re.sub(r"!\[(.*?)\]\(.*?\)", r"\1", t)          # images -> alt
    t = re.sub(r"\[(.*?)\]\(.*?\)", r"\1", t)            # links -> text
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)               # bold
    t = re.sub(r"(?m)^#{1,6}\s*", "", t)                 # heading markers at line start
    t = re.sub(r"#{2,}\s*", "", t)                       # leftover ## / ### mid-line
    t = re.sub(r"(?<=\s)#(?=\s)", "", t)                 # lone # mid-line
    t = t.replace("`", "")
    t = re.sub(r"\s+", " ", t).strip()
    return t


def clean_query_for_header(query: str, limit: int = 220) -> str:
    """One clean single-line topic for the PDF/MD header — never a raw prompt dump."""
    t = clean_inline(query)
    if len(t) <= limit:
        return t
    cut = t[:limit].rsplit(" ", 1)[0]
    return (cut or t[:limit]).rstrip(" ,;:") + "…"


def normalize_bullets(md_text: str) -> str:
    """Turn inline '- a - b - c' runs inside a paragraph into real markdown list blocks.

    LLM output often writes: 'Key elements include: - Legal accountability: ... - Physical dexterity: ...'
    on ONE line. Markdown libs keep that as literal '- ' text inside a <p>. Splitting them
    into separate '- ' lines (with a blank line before) makes them real lists in PDF and DOCX.
    """
    def _is_list_item(s: str) -> bool:
        t = s.strip()
        return (t.startswith(("-", "*")) and re.match(r"^[-*]\s+\S", t) is not None) \
            or re.match(r"^\d+[.)]\s+\S", t) is not None

    out = []
    prev_was_list = True  # start-of-doc needs no blank line
    for line in (md_text or "").split("\n"):
        st = line.strip()
        if _is_list_item(line):
            # markdown libs only open a real <ul>/<ol> after a blank line —
            # without it, '- ' lines merge into the paragraph above as literal text
            if not prev_was_list:
                out.append("")
            out.append(line)
            prev_was_list = True
            continue
        prev_was_list = (not st)
        # already structural lines: leave untouched
        if (not st or st.startswith((">", "#", "|", "```"))):
            out.append(line)
            continue
        if " - " not in line:
            out.append(line)
            continue
        parts = re.split(r"\s+-\s+(?=[A-Z0-9\"'])", line)
        if len(parts) < 2:
            out.append(line)
            continue
        # only split when it really looks like a bullet run, not prose with dashes
        first_ends_list_intro = parts[0].rstrip().endswith((":", ";"))
        if not first_ends_list_intro and len(parts) < 3:
            out.append(line)
            continue
        out.append(parts[0].rstrip())
        out.append("")  # blank line so markdown libs open a real list
        for p in parts[1:]:
            out.append("- " + p.strip())
    return "\n".join(out)


def _shorten(text: str, limit: int) -> str:
    t = (text or "").strip()
    if len(t) <= limit:
        return t
    cut = t[:limit].rsplit(" ", 1)[0]
    return (cut or t[:limit]).rstrip(" ,;:") + "…"


def _is_table_sep(line: str) -> bool:
    s = (line or "").strip()
    return s.startswith("|") and "-" in s and set(s) <= set("|: -")


def normalize_tables(md_text: str) -> str:
    """Ensure a blank line before a markdown table block.

    Markdown table extensions (python-markdown AND remark-gfm) only open a
    real <table> when the header row is preceded by a blank line. LLM output
    often writes '**Comparison:**\\n| Program |...' with no gap, which then
    renders as a paragraph of raw pipes ('Comparison: | Program | ... ||
    Penn Foster | ...'). Insert the missing blank line before the header row.
    Continuation rows need no blank line.
    """
    lines = (md_text or "").split("\n")
    out = []
    prev_blank_or_table = True  # start-of-doc needs no blank line
    for i, line in enumerate(lines):
        st = line.strip()
        is_pipe = st.startswith("|")
        if is_pipe and not prev_blank_or_table:
            nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
            if _is_table_sep(nxt):
                out.append("")
        out.append(line)
        prev_blank_or_table = (not st) or is_pipe
    return "\n".join(out)


def normalize_markdown(md_text: str) -> str:
    """All render-safety normalizations in one place (bullets + tables)."""
    return normalize_tables(normalize_bullets(md_text or ""))


def _norm_title(t: str) -> str:
    return re.sub(r"\s+", " ", (t or "").strip().lower().rstrip(" +"))


def _truncate_html_cell(cell_html: str, limit: int) -> str:
    """Cap a table cell's visible text; over-long cells fall back to plain text."""
    plain = _strip_tags(cell_html)
    if len(plain) <= limit:
        return cell_html
    return _escape(_shorten(plain, limit))


def build_docx(title: str, report_md: str, query: str = "") -> io.BytesIO:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement

    def _repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)

    def _add_hyperlink(paragraph, url: str, display: str, size=11):
        """Clickable blue hyperlink run (python-docx has no native API)."""
        part = paragraph.part
        r_id = part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(_qn("r:id"), r_id)
        run_el = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(_qn("w:val"), str(int(size * 2)))
        rPr.append(sz)
        col = OxmlElement("w:color")
        col.set(_qn("w:val"), "0563C1")
        rPr.append(col)
        u = OxmlElement("w:u")
        u.set(_qn("w:val"), "single")
        rPr.append(u)
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(_qn("w:ascii"), "Calibri")
        rFonts.set(_qn("w:hAnsi"), "Calibri")
        rPr.append(rFonts)
        run_el.append(rPr)
        t_el = OxmlElement("w:t")
        t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t_el.text = display
        run_el.append(t_el)
        hyperlink.append(run_el)
        paragraph._p.append(hyperlink)

    _LINK_RE = re.compile(r"(\[[^\]]+\]\([^)]+\)|https?://[^\s<>()]+)")

    def _rich_text(paragraph, text: str, size=11, bold=False, italic=False, color=None):
        # links first ([text](url) + bare URLs) -> clickable runs, then
        # minimal **bold** / *italic* / `code` support inside a paragraph
        for seg in _LINK_RE.split(text):
            if not seg:
                continue
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)$", seg)
            bare = re.match(r"(https?://[^\s<>()]+)$", seg)
            if m and m.group(2).startswith("http"):
                _add_hyperlink(paragraph, m.group(2), m.group(1), size=size)
                continue
            if bare:
                url = bare.group(1).rstrip(".,;:")
                trail = bare.group(1)[len(url):]
                _add_hyperlink(paragraph, url, url, size=size)
                if trail:
                    r = paragraph.add_run(trail)
                    r.font.size = Pt(size)
                continue
            tokens = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`.+?`)", seg)
            for tok in tokens:
                if not tok:
                    continue
                run = paragraph.add_run()
                if tok.startswith("**") and tok.endswith("**"):
                    run.text = tok[2:-2]
                    run.bold = True
                elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                    run.text = tok[1:-1]
                    run.italic = True
                elif tok.startswith("`") and tok.endswith("`"):
                    run.text = tok[1:-1]
                    run.font.name = "Consolas"
                else:
                    run.text = re.sub(r"\[(.*?)\]\(.*?\)", r"\1", tok)
                run.font.size = Pt(size)
                if bold:
                    run.bold = True
                if italic:
                    run.italic = True
                if color is not None:
                    run.font.color.rgb = color
        return paragraph

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(clean_inline(title) or "Research Report", level=0)
    # NOTE: no Query/prompt block here — the export is the report, not the input prompt.

    body = normalize_markdown(report_md or "")
    lines = body.split("\n")
    # drop a leading H1 that duplicates the cover title
    if lines and lines[0].strip().startswith("# "):
        h1 = clean_inline(lines[0].strip()[2:])
        if _norm_title(h1) == _norm_title(title):
            lines = lines[1:]

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code_lines))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue
        if stripped.startswith("#### "):
            doc.add_heading(clean_inline(stripped[5:]), level=4)
        elif stripped.startswith("### "):
            doc.add_heading(clean_inline(stripped[4:]), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(clean_inline(stripped[3:]), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(clean_inline(stripped[2:]), level=1)
        elif re.match(r"^(\-|\*) ", stripped):
            text = re.sub(r"^(\-|\*) ", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            _rich_text(p, text)
        elif re.match(r"^\d+[.)] ", stripped):
            text = re.sub(r"^\d+[.)] ", "", stripped)
            p = doc.add_paragraph(style="List Number")
            _rich_text(p, text)
        elif stripped.startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in tbl_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(set(c) <= set("-: ") for c in cells):
                    continue  # separator row
                rows.append(cells)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Light Grid Accent 1"
                table.autofit = True
                _repeat_header(table.rows[0])  # repeat header on page breaks
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        val = row[ci] if ci < len(row) else ""
                        cell = table.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        if ri == 0:
                            _rich_text(p, val, size=10, bold=True,
                                       color=RGBColor(0xFF, 0xFF, 0xFF))
                            shading = OxmlElement("w:shd")
                            shading.set(_qn("w:fill"), "4338CA")
                            p._p.get_or_add_pPr().append(shading)
                        else:
                            _rich_text(p, val, size=10)
            continue
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            _rich_text(p, clean_inline(stripped.lstrip("> ").strip()),
                       size=10, italic=True, color=RGBColor(0x55, 0x55, 0x55))
        else:
            p = doc.add_paragraph()
            _rich_text(p, stripped)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _qn(prefix_attr: str) -> str:
    from docx.oxml.ns import qn
    return qn(prefix_attr)


def build_pdf(title: str, report_md: str, query: str = "") -> io.BytesIO:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, ListFlowable, ListItem,
    )
    from reportlab.lib.enums import TA_LEFT
    import markdown as md_lib

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title=clean_inline(title) or "Research Report",
    )
    avail = doc.width
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=20, leading=24, spaceAfter=6))
    styles.add(ParagraphStyle("H2Brand", parent=styles["Heading2"], textColor=HexColor("#4338ca"), fontSize=14, leading=18, spaceBefore=14, spaceAfter=6))
    styles.add(ParagraphStyle("BodyPro", parent=styles["Normal"], fontSize=10.5, leading=15, alignment=TA_LEFT, spaceAfter=5))
    styles.add(ParagraphStyle("QuotePro", parent=styles["Normal"], fontSize=10, leading=14, textColor=HexColor("#555555"), leftIndent=12, borderPadding=(6, 6, 6)))

    clean_title = clean_inline(title) or "Research Report"
    story = [Paragraph(_escape(clean_title), styles["ReportTitle"])]
    short_q = clean_query_for_header(query)
    if short_q:
        story.append(Paragraph(f"<i>Topic: {_escape(short_q)}</i>", styles["BodyPro"]))
    story.append(HRFlowable(width="100%", thickness=1, color=HexColor("#e4e4e7"), spaceAfter=8))

    body = normalize_markdown(report_md or "")
    # drop a leading H1 that duplicates the cover title
    lines_tmp = body.split("\n")
    if lines_tmp and lines_tmp[0].strip().startswith("# "):
        h1 = clean_inline(lines_tmp[0].strip()[2:])
        if _norm_title(h1) == _norm_title(clean_title):
            body = "\n".join(lines_tmp[1:])

    html = md_lib.markdown(body, extensions=["tables", "fenced_code"])
    blocks = re.split(r"(?=<h[1-4]|<ul|<ol|<table|<blockquote|<pre|<hr|<p)", html)
    for b in blocks:
        b = b.strip()
        if not b:
            continue
        if b.startswith("<h1"):
            story.append(Paragraph(_escape(_strip_tags(b)), styles["Heading1"]))
        elif b.startswith("<h2"):
            story.append(Paragraph(_escape(_strip_tags(b)), styles["H2Brand"]))
        elif b.startswith("<h3") or b.startswith("<h4"):
            story.append(Paragraph(_escape(_strip_tags(b)), styles["Heading3"]))
        elif b.startswith("<table"):
            story.extend(_pdf_table(b, avail))
        elif b.startswith("<ul") or b.startswith("<ol"):
            items = re.findall(r"<li.*?>(.*?)</li>", b, flags=re.S)
            lis = [ListItem(Paragraph(_inline(x), styles["BodyPro"])) for x in items[:40]]
            if lis:
                story.append(ListFlowable(lis, bulletType="bullet" if b.startswith("<ul") else "1", leftIndent=18))
                story.append(Spacer(1, 4))
        elif b.startswith("<blockquote"):
            story.append(Paragraph(_inline(b), styles["QuotePro"]))
            story.append(Spacer(1, 6))
        elif b.startswith("<pre"):
            code = _strip_tags(b)[:2000]
            story.append(Paragraph(f'<font face="Courier" size="8">{_escape(code)}</font>', styles["BodyPro"]))
        elif b.startswith("<hr"):
            story.append(HRFlowable(width="100%", thickness=0.7, color=HexColor("#e4e4e7"), spaceBefore=8, spaceAfter=8))
        elif b.startswith("<p"):
            txt = _inline(b)
            if txt.strip():
                story.append(Paragraph(txt, styles["BodyPro"]))

    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=0.7, color=HexColor("#e4e4e7")))
    story.append(Paragraph("Generated by lianResearch • agent-reach deep research", styles["BodyPro"]))

    doc.build(story)
    buf.seek(0)
    return buf


def _pdf_table(html_block: str, avail_width: float):
    """Build a reportlab Table with wrapping Paragraph cells and content-aware widths.

    Root fix for the overlap bug: cells were raw strings (no word wrap) in equal
    fixed-width columns, so long text painted over neighbouring cells and wide
    tables ran off the page edge.
    """
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor

    plain_rows = []
    rich_rows = []
    for tr in re.findall(r"<tr.*?>(.*?)</tr>", html_block, flags=re.S):
        cells = re.findall(r"<t[hd].*?>(.*?)</t[hd]>", tr, flags=re.S)
        plain_rows.append([_shorten(_strip_tags(c), 500) for c in cells])
        rich_rows.append([_inline(_truncate_html_cell(c, 500)) for c in cells])
    if not plain_rows:
        return []
    raw_rows = plain_rows  # plain text used for width measuring
    ncols = max(len(r) for r in raw_rows)
    # content-proportional weights (header counts double so headers don't crush)
    weights = []
    for c in range(ncols):
        w = 0
        for ri, r in enumerate(raw_rows):
            cell = r[c] if c < len(r) else ""
            w += len(cell) * (2 if ri == 0 else 1)
        weights.append(max(w, 10))
    total = sum(weights)
    widths = [avail_width * w / total for w in weights]
    # enforce a minimum readable width (fits ~'Switzerland' at 7.5pt), then rescale
    widths = [max(x, 54) for x in widths]
    over = sum(widths) / avail_width
    if over > 1:
        widths = [x / over for x in widths]

    if ncols >= 6:
        font_size, leading, pad = 7.5, 10, 3
    elif ncols >= 5:
        font_size, leading, pad = 8, 11, 4
    else:
        font_size, leading, pad = 8.5, 12, 5

    head_style = ParagraphStyle("CellHeadX", fontSize=font_size, leading=leading,
                                textColor=HexColor("#ffffff"), fontName="Helvetica-Bold")
    cell_style = ParagraphStyle("CellProX", fontSize=font_size, leading=leading,
                                fontName="Helvetica")
    data = []
    for ri in range(len(plain_rows)):
        row = []
        for ci in range(ncols):
            rich = rich_rows[ri][ci] if ci < len(rich_rows[ri]) else ""
            row.append(Paragraph(rich, head_style if ri == 0 else cell_style))
        data.append(row)

    t = Table(data, colWidths=widths, repeatRows=1, splitByRow=True)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), HexColor("#4338ca")),
        ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#ffffff")),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#d4d4d8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#ffffff"), HexColor("#f4f4f5")]),
        ("TOPPADDING", (0, 0), (-1, -1), pad),
        ("BOTTOMPADDING", (0, 0), (-1, -1), pad),
        ("LEFTPADDING", (0, 0), (-1, -1), pad),
        ("RIGHTPADDING", (0, 0), (-1, -1), pad),
    ]))
    return [t, Spacer(1, 8)]


_EMOJI_RE = re.compile(
    "[\U0001F300-\U0001FAFF"  # symbols, pictographs, emoticons, transport
    "\u2600-\u27BF"            # misc symbols, dingbats
    "\u2B00-\u2BFF"            # misc symbols and arrows
    "\uFE00-\uFE0F"            # variation selectors
    "\U0001F1E6-\U0001F1FF"    # regional indicators (flags)
    "]+",
    flags=re.UNICODE,
)


def _strip_tags(h: str) -> str:
    import html as _html
    t = re.sub(r"<.*?>", "", h, flags=re.S)
    t = _EMOJI_RE.sub("", t)  # Helvetica (WinAnsi) has no emoji glyphs -> would render as boxes
    return _html.unescape(t).strip()


def _escape(t: str) -> str:
    return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# reportlab paragraph tags we allow through (everything else is stripped)
_INLINE_ALLOWED = {"b", "i", "u", "font", "link", "br", "sub", "sup"}


def _sanitize_tag(tag: str) -> str:
    """Keep only safe reportlab tags/attrs; drop the rest (e.g. <a>, <span>, <div>)."""
    m = re.match(r"</?([a-zA-Z][a-zA-Z0-9]*)", tag)
    if not m or m.group(1).lower() not in _INLINE_ALLOWED:
        return ""
    name = m.group(1).lower()
    if tag.startswith("</"):
        return f"</{name}>"
    if name == "link":
        href = re.search(r'href="([^"]+)"', tag)
        url = html_unescape(href.group(1)) if href else ""
        if not re.match(r"^https?://", url):
            return ""
        return f'<link href="{_escape(url)}" color="blue">'
    if name == "font":
        attrs = []
        face = re.search(r'face="([^"]+)"', tag)
        if face and face.group(1).lower() in ("courier", "courier-bold"):
            attrs.append(f'face="{face.group(1)}"')
        color = re.search(r'color="([^"]+)"', tag)
        if color and re.match(r"^(blue|red|green|black|#[0-9a-fA-F]{6})$", color.group(1)):
            attrs.append(f'color="{color.group(1)}"')
        return f"<font {' '.join(attrs)}>" if attrs else "<font>"
    if name == "br":
        return "<br/>"
    return f"<{name}>"


def _escape_keep_tags(s: str) -> str:
    """Escape text nodes but preserve allowed reportlab inline tags."""
    parts = re.split(r"(<[^>]+>)", s)
    out = []
    for p in parts:
        if p.startswith("<") and p.endswith(">"):
            out.append(_sanitize_tag(p))
        else:
            out.append(_escape(html_unescape(p)))
    return "".join(out)


def html_unescape(t: str) -> str:
    import html as _html
    return _html.unescape(t or "")


def _inline(html_block: str) -> str:
    """Convert an HTML fragment to reportlab paragraph markup.

    Keeps <b>/<i>/code and turns <a> into real clickable blue links
    instead of flattening everything to plain text.
    """
    s = html_block or ""
    s = re.sub(r"<strong.*?>(.*?)</strong>", r"<b>\1</b>", s, flags=re.S)
    s = re.sub(r"<em.*?>(.*?)</em>", r"<i>\1</i>", s, flags=re.S)
    s = re.sub(r"<code.*?>(.*?)</code>", r"<font face=\"Courier\">\1</font>", s, flags=re.S)

    def _link_repl(m):
        url = html_unescape(m.group(1)).strip()
        text = m.group(2).strip() or url
        if not re.match(r"^https?://", url):
            return text
        return f'<link href="{url}"><u><font color="blue">{text}</font></u></link>'

    s = re.sub(r'<a.*?href="(.*?)".*?>(.*?)</a>', _link_repl, s, flags=re.S)
    # bare URLs that markdown left as text -> clickable too
    s = re.sub(r"(?<!href=\")(?<!>)(https?://[^\s<>\"']+)",
               lambda m: f'<link href="{m.group(1)}"><u><font color="blue">{m.group(1)}</font></u></link>', s)
    s = _EMOJI_RE.sub("", s)
    return _escape_keep_tags(s)
